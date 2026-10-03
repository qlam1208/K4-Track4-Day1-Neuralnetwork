"""results_table.py — Lưu kết quả thí nghiệm ra JSON và xuất bảng Excel experiments.xlsx.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx.
"""
from __future__ import annotations

import json
from pathlib import Path
import openpyxl

FORMULA_COLS = {
    "step0_gap_vs_lnC",
    "gap_val_minus_train",
    "delta_val_f1_vs_base",
    "beyond_noise",
}


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file.
    """
    p_dir = Path(results_dir)
    p_dir.mkdir(parents=True, exist_ok=True)
    exp_id = result["cfg"].get("exp_id", "exp")
    out_path = p_dir / f"{exp_id}.json"

    payload = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"],
    }
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2)

    return str(out_path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    p_dir = Path(results_dir)
    if not p_dir.exists():
        return []

    results = []
    for fp in p_dir.glob("*.json"):
        with open(fp, "r", encoding="utf-8") as f:
            data = json.load(f)
            results.append(data)

    results.sort(key=lambda r: r.get("cfg", {}).get("exp_id", ""))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png".
    """
    cfg = result["cfg"]
    summary = result["summary"]
    exp_id = cfg.get("exp_id", "")

    hidden = cfg.get("hidden", "")
    if isinstance(hidden, (list, tuple)):
        hidden_str = "-".join(map(str, hidden))
    else:
        hidden_str = str(hidden)

    row = {
        "exp_id": exp_id,
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce"),
        "optimizer": cfg.get("optimizer", ""),
        "lr": cfg.get("lr", ""),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": hidden_str,
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": cfg.get("clip_norm", "") if cfg.get("clip_norm") is not None else "None",
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 1),
        "step0_loss": summary.get("step0_loss", ""),
        "best_val_loss": summary.get("best_val_loss", ""),
        "best_epoch": summary.get("best_epoch", ""),
        "final_train_loss": summary.get("final_train_loss", ""),
        "final_val_loss": summary.get("final_val_loss", ""),
        "val_acc": summary.get("val_acc", ""),
        "val_macro_f1": summary.get("val_macro_f1", ""),
        "time_per_epoch_s": summary.get("time_per_epoch_s", ""),
        "peak_mem_MB": summary.get("peak_mem_MB", ""),
        "diverged": summary.get("diverged", False),
        "eval_acc": eval_scores.get("acc", "") if eval_scores else "",
        "eval_macro_f1": eval_scores.get("macro_f1", "") if eval_scores else "",
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes,
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path."""
    p_out = Path(out_path)
    p_out.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    headers = [cell.value for cell in ws[1]]

    for i, row in enumerate(rows):
        row_idx = 2 + i
        for col_idx, h in enumerate(headers, start=1):
            if h is None or h in FORMULA_COLS:
                continue
            if h in row:
                ws.cell(row=row_idx, column=col_idx, value=row[h])

    wb.save(p_out)
    print(f"Đã lưu bảng thí nghiệm ({len(rows)} dòng) ra {out_path}")
